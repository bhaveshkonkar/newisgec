"""Office/text extraction regression.

Builds real DOCX/XLSX packages rather than mocking them, because the whole
point of these tests is that the bytes on disk are what a browser would send.
A hand-written fake would pass even if the ZIP layout were wrong.

Run: .\\.venv\\Scripts\\python.exe scripts\\test_docparse.py
"""
from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))

import docparse  # noqa: E402

TMP = Path(__file__).resolve().parent.parent / "data" / "_docparse_fixtures"

PASSED: list[str] = []
FAILED: list[tuple[str, str]] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    if cond:
        PASSED.append(name)
        print(f"  PASS  {name}")
    else:
        FAILED.append((name, detail))
        print(f"  FAIL  {name}" + (f"\n          {detail}" if detail else ""))


def section(t: str) -> None:
    print(f"\n{t}\n" + "-" * len(t))


# --------------------------------------------------------------- fixtures

CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
</Types>"""


def make_docx(path: Path, paragraphs: list[str], table: list[list[str]] | None = None) -> None:
    body = "".join(
        f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in paragraphs
    )
    if table:
        rows = "".join(
            "<w:tr>"
            + "".join(f"<w:tc><w:p><w:r><w:t>{c}</w:t></w:r></w:p></w:tc>" for c in row)
            + "</w:tr>"
            for row in table
        )
        body += f"<w:tbl>{rows}</w:tbl>"
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{docparse.W_NS}"><w:body>{body}</w:body></w:document>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", CONTENT_TYPES)
        zf.writestr("word/document.xml", document)


def make_xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Selection Table"
    ws.append(["Series", "Size", "Power kW", "Ratio", "Service factor", "Weight kg"])
    ws.append(["S3", "S3-40", "3.0", "17", "1.25", "18.5"])
    ws.append(["S3", "S3-50", "4.0", "19", "1.30", "22.0"])
    ws2 = wb.create_sheet("Hidden")
    ws2.append(["should not appear"])
    ws2.sheet_state = "hidden"
    wb.save(path)


def main() -> int:
    TMP.mkdir(parents=True, exist_ok=True)

    section("DOCX: paragraphs and tables keep document order")
    d = TMP / "sample.docx"
    make_docx(d, [
        "Elecon G B_EON Series gearbox catalogue",
        "Service factor 1.25 at 1500 rpm input speed.",
        "Mounting: foot and torque arm available.",
    ], table=[
        ["Series", "Size", "Power kW", "Ratio"],
        ["S3", "S3-40", "3.0", "17"],
        ["S3", "S3-50", "4.0", "19"],
    ])
    text, pages, meta = docparse.extract(d, 100_000)
    check("kind is docx", meta["kind"] == "docx", str(meta))
    check("page_count is None, not invented", pages is None, str(pages))
    check("paragraphs extracted", "Service factor 1.25 at 1500 rpm" in text, text[:200])
    check("no zip noise in the output", "PK" not in text[:200], text[:80])
    check("no XML tags leaked", "<w:" not in text, text[:200])
    for expect in ("S3-40", "S3-50", "Series"):
        check(f"table cell {expect!r} present", expect in text)
    check("table row separator used", "S3 | S3-40 | 3.0 | 17" in text, text)
    i_par = text.index("Service factor")
    i_tbl = text.index("Series | Size")
    check("table appears after its paragraph, in order", i_par < i_tbl, f"{i_par} vs {i_tbl}")

    section("DOCX: a spreadsheet renamed to .docx is refused, not garbled")
    x = TMP / "actually_sheet.docx"
    make_xlsx(x)
    try:
        docparse.extract(x, 100_000)
        check("renamed xlsx as docx is refused", False, "no exception raised")
    except docparse.UnsupportedDocument as e:
        check("renamed xlsx as docx is refused", "spreadsheet" in str(e).lower(), str(e))

    section("DOCX: corrupt zip is a clean 422, not a traceback")
    bad = TMP / "corrupt.docx"
    bad.write_bytes(b"PK\x03\x04" + b"\x00" * 200)
    try:
        docparse.extract(bad, 100_000)
        check("corrupt docx raises UnsupportedDocument", False, "no exception")
    except docparse.UnsupportedDocument as e:
        check("corrupt docx raises UnsupportedDocument", "Office file" in str(e), str(e))

    section("XLSX: values, sheet names, hidden sheets skipped")
    s = TMP / "sample.xlsx"
    make_xlsx(s)
    text, pages, meta = docparse.extract(s, 100_000)
    check("kind is xlsx", meta["kind"] == "xlsx", str(meta))
    check("sheet count reported", meta.get("sheets") == 1, str(meta))
    check("visible sheet is named in the output", "# sheet: Selection Table" in text, text[:120])
    check("header row present", "Service factor" in text, text[:200])
    for expect in ("S3-40", "3.0", "1.25", "18.5"):
        check(f"cell value {expect!r} present", expect in text)
    check("hidden sheet excluded", "should not appear" not in text, text)

    section("XLSX: numeric cells are not returned as datetimes/percentages")
    pct = TMP / "types.xlsx"
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["kw", "sf", "ratio", "rpm"])
    ws.append([60, 1.25, 17, 1500])
    wb.save(pct)
    text, _, _ = docparse.extract(pct, 100_000)
    row = [ln for ln in text.splitlines() if "60" in ln]
    check("numbers come out as plain numbers", row and "60\t1.25\t17\t1500" in text, text)

    section("Legacy .doc/.xls are refused with an actionable message")
    ole = TMP / "old.xls"
    ole.write_bytes(docparse.OLE2_MAGIC + b"\x00" * 512)
    for suffix in (".doc", ".xls"):
        f = TMP / f"old{suffix}"
        f.write_bytes(docparse.OLE2_MAGIC + b"\x00" * 512)
        try:
            docparse.extract(f, 100_000)
            check(f"{suffix} refused", False, "no exception")
        except docparse.UnsupportedDocument as e:
            check(f"{suffix} refused", "Save As" in str(e), str(e))
            check(f"{suffix} message names the modern target",
                  (".docx" if suffix == ".doc" else ".xlsx") in str(e), str(e))

    section("CSV/TXT: encoding is sniffed, not assumed")
    utf8 = TMP / "u8.csv"
    utf8.write_bytes("Series,Power kW\nS3,3.0\n".encode("utf-8"))
    t, _, m = docparse.extract(utf8, 100_000)
    check("utf-8 csv read", "S3,3.0" in t, t)

    bom = TMP / "bom.csv"
    bom.write_bytes(b"\xef\xbb\xbf" + "Series,Power kW\nS3,3.0\n".encode("utf-8"))
    t, _, _ = docparse.extract(bom, 100_000)
    check("utf-8 BOM stripped, not rendered", not t.startswith("\ufeff"), repr(t[:10]))

    u16 = TMP / "u16.csv"
    u16.write_bytes("Series,Power kW\nS3,3.0\n".encode("utf-16"))
    t, _, _ = docparse.extract(u16, 100_000)
    check("utf-16 with BOM decoded", "S3,3.0" in t, repr(t[:40]))

    u16nb = TMP / "u16nb.csv"
    u16nb.write_bytes("Series,Power kW\nS3,3.0\n".encode("utf-16-le"))
    t, _, _ = docparse.extract(u16nb, 100_000)
    check("utf-16 without BOM detected", "S3,3.0" in t, repr(t[:60]))

    section("Unknown extension: renamed Office files still parse")
    r = TMP / "catalogue.dat"
    r.write_bytes(x.read_bytes())
    t, _, m = docparse.extract(r, 100_000)
    check("renamed xlsx parsed by content", m["kind"] == "xlsx", str(m))
    check("rename recorded in metadata", m.get("renamed_from") == ".dat", str(m))

    binf = TMP / "thing.dat"
    binf.write_bytes(b"\x00\x01\x02" * 500)
    try:
        docparse.extract(binf, 100_000)
        check("unrecognised binary refused", False, "no exception")
    except docparse.UnsupportedDocument as e:
        check("unrecognised binary refused", "binary" in str(e), str(e))

    section("Bounded output")
    big = TMP / "big.docx"
    make_docx(big, ["service factor 1.25 rpm 1500 kW gearbox"] * 5000)
    t, _, m = docparse.extract(big, 2_000)
    check("output respects max_chars", len(t) <= 2_000, str(len(t)))
    check("truncation is honest, not a crash", t.startswith("service factor"), t[:60])

    section("suspect-text guard")
    check("real catalogue text passes the guard",
          docparse.looks_like_text("Elecon gearbox series S3, 3.0 kW at 1500 rpm, ratio 17, service factor 1.25"))
    check("mojibake fails the guard",
          not docparse.looks_like_text("PK\x03\x04�����������"))
    check("empty text fails the guard", not docparse.looks_like_text("   "))

    print("\n" + "=" * 60)
    print(f"PASSED {len(PASSED)}   FAILED {len(FAILED)}")
    for name, detail in FAILED:
        print(f"  - {name}" + (f"\n      {detail}" if detail else ""))
    print("=" * 60)

    import shutil

    shutil.rmtree(TMP, ignore_errors=True)
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main())